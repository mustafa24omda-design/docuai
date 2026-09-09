"""
DocuAI - Intelligent Project Document Management System
=========================================================
نسخة مدمجة بالكامل في ملف واحد (app.py) — بدون الحاجة لمجلدات فرعية،
لتسهيل الرفع عبر واجهة GitHub الويب والنشر المباشر على Render.

يحتوي هذا الملف على:
1. محرك الاستخراج (Extraction Engine) بالـ Regex.
2. قارئات الملفات (PDF / DOCX / TXT).
3. واجهة ويب (Flask) لرفع المستندات ومعاينة/تنزيل النتائج.
"""

import os
import re
import socket
import tempfile
import threading
import webbrowser
from pathlib import Path
from datetime import datetime
from dataclasses import dataclass, asdict
from typing import Optional

from flask import Flask, request, render_template_string, send_file


# =============================================================================
# 1) محرك الاستخراج (Extraction Engine)
# =============================================================================

DOC_NUMBER_PATTERNS = [
    r"(?:Document\s*(?:No\.?|Number|#)\s*[:\-]?\s*)([A-Z0-9][A-Z0-9\-\/\.]{3,30})",
    r"(?:Doc\.?\s*(?:No\.?|#)\s*[:\-]?\s*)([A-Z0-9][A-Z0-9\-\/\.]{3,30})",
    r"(?:Ref(?:erence)?\.?\s*(?:No\.?)?\s*[:\-]?\s*)([A-Z0-9][A-Z0-9\-\/\.]{3,30})",
    r"\b([A-Z]{2,6}-[A-Z0-9]{2,10}-\d{2,6}(?:-[A-Z0-9]{1,4})?)\b",
]

DOC_TYPE_KEYWORDS = {
    "Transmittal": ["transmittal", "خطاب إحالة", "إحالة مستند"],
    "RFI": ["request for information", "rfi", "طلب معلومات"],
    "Submittal": ["submittal", "material submittal", "shop drawing submittal"],
    "Shop Drawing": ["shop drawing", "مخطط تنفيذي"],
    "Method Statement": ["method statement", "منهجية تنفيذ"],
    "Inspection Request": ["inspection request", "ir ", "طلب معاينة"],
    "Non-Conformance Report": ["non-conformance", "ncr", "تقرير عدم مطابقة"],
    "Minutes of Meeting": ["minutes of meeting", "mom", "محضر اجتماع"],
    "Design Drawing": ["design drawing", "مخطط تصميمي"],
    "Correspondence": ["letter", "correspondence", "خطاب", "مراسلة"],
    "Progress Report": ["progress report", "تقرير سير العمل"],
    "Invoice": ["invoice", "فاتورة"],
}

DATE_PATTERNS = [
    r"\b(\d{4}[-/]\d{1,2}[-/]\d{1,2})\b",
    r"\b(\d{1,2}[-/]\d{1,2}[-/]\d{4})\b",
    r"\b(\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4})\b",
    r"\b([A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4})\b",
]

DATE_FORMATS = [
    "%Y-%m-%d", "%Y/%m/%d",
    "%d-%m-%Y", "%d/%m/%Y",
    "%d %B %Y", "%d %b %Y",
    "%B %d, %Y", "%b %d, %Y", "%B %d %Y", "%b %d %Y",
]

REVISION_PATTERNS = [
    r"(?:Revision|Rev\.?)\s*[:\-]?\s*([A-Z]?\d{1,3}|[A-Z])\b",
    r"(?:مراجعة|مراجعه|النسخة)\s*[:\-]?\s*(\d{1,3})",
]

STATUS_KEYWORDS = {
    "Approved": ["approved", "معتمد"],
    "Approved as Noted": ["approved as noted", "معتمد مع ملاحظات"],
    "Rejected": ["rejected", "not approved", "مرفوض"],
    "Revise and Resubmit": ["revise and resubmit", "أعد التقديم"],
    "For Review": ["for review", "under review", "قيد المراجعة"],
    "For Approval": ["for approval", "للاعتماد"],
    "For Information": ["for information", "for your information", "للعلم"],
    "Closed": ["closed", "مغلق"],
    "Open": ["open", "مفتوح"],
    "Pending": ["pending", "معلق"],
}

SENDER_PATTERNS = [
    r"(?:From|Sender|Originator|Issued\s*By)\s*[:\-]\s*([A-Za-z0-9&.,\-\s]{2,60})(?:\n|$)",
    r"(?:من|المرسل|جهة الإصدار)\s*[:\-]\s*([\u0600-\u06FF0-9A-Za-z&.,\-\s]{2,60})(?:\n|$)",
]


@dataclass
class ExtractedDocument:
    source_file: str
    document_number: Optional[str] = None
    document_type: Optional[str] = None
    date: Optional[str] = None
    revision: Optional[str] = None
    status: Optional[str] = None
    sender: Optional[str] = None
    confidence: Optional[str] = None

    def to_dict(self):
        return asdict(self)


def _first_match(patterns, text, flags=re.IGNORECASE):
    for pat in patterns:
        m = re.search(pat, text, flags)
        if m:
            return m.group(1).strip().rstrip(".,;")
    return None


def _extract_document_type(text: str) -> Optional[str]:
    lowered = text.lower()
    for doc_type, keywords in DOC_TYPE_KEYWORDS.items():
        for kw in keywords:
            if kw.lower() in lowered:
                return doc_type
    return None


def _normalize_date(raw_date: str) -> Optional[str]:
    if not raw_date:
        return None
    cleaned = raw_date.strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(cleaned, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return cleaned


def _match_status_in(snippet: str) -> Optional[str]:
    lowered = snippet.lower()
    all_kw = [(status, kw) for status, kws in STATUS_KEYWORDS.items() for kw in kws]
    all_kw.sort(key=lambda x: len(x[1]), reverse=True)
    for status, kw in all_kw:
        if kw.lower() in lowered:
            return status
    return None


def _extract_status(text: str) -> Optional[str]:
    label_match = re.search(r"(?:Status|الحالة)\s*[:\-]\s*(.+)", text, re.IGNORECASE)
    if label_match:
        found = _match_status_in(label_match.group(1))
        if found:
            return found
    return _match_status_in(text)


def extract_fields(text: str, source_file: str = "") -> ExtractedDocument:
    doc_number = _first_match(DOC_NUMBER_PATTERNS, text)
    doc_type = _extract_document_type(text)
    raw_date = _first_match(DATE_PATTERNS, text)
    date = _normalize_date(raw_date) if raw_date else None
    revision = _first_match(REVISION_PATTERNS, text)
    status = _extract_status(text)
    sender = _first_match(SENDER_PATTERNS, text)

    found_count = sum(1 for v in [doc_number, doc_type, date, revision, status, sender] if v)
    if found_count >= 5:
        confidence = "High"
    elif found_count >= 3:
        confidence = "Medium"
    else:
        confidence = "Low"

    return ExtractedDocument(
        source_file=source_file,
        document_number=doc_number,
        document_type=doc_type,
        date=date,
        revision=revision,
        status=status,
        sender=sender,
        confidence=confidence,
    )


# =============================================================================
# 2) قراءة الملفات (PDF / DOCX / TXT)
# =============================================================================

def read_pdf(path: str) -> str:
    import pdfplumber
    text_parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
    return "\n".join(text_parts)


def read_docx(path: str) -> str:
    import docx
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def read_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


READERS = {".pdf": read_pdf, ".docx": read_docx, ".txt": read_txt}


def read_any(path: str) -> str:
    ext = Path(path).suffix.lower()
    reader = READERS.get(ext)
    if reader is None:
        raise ValueError(f"صيغة الملف غير مدعومة: {ext}")
    return reader(path)


# =============================================================================
# 3) تصدير النتائج إلى Excel
# =============================================================================

def export_to_excel(results, output_path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "DocuAI Results"

    headers = [
        "اسم الملف", "رقم المستند", "نوع المستند",
        "التاريخ", "Revision", "Status", "الجهة المرسلة", "دقة الاستخراج",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in results:
        ws.append([
            r.source_file, r.document_number, r.document_type,
            r.date, r.revision, r.status, r.sender, r.confidence,
        ])

    for col_idx, header in enumerate(headers, start=1):
        max_len = max(
            [len(str(header))]
            + [len(str(ws.cell(row=row, column=col_idx).value or "")) for row in range(2, ws.max_row + 1)]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 45)

    ws.freeze_panes = "A2"
    wb.save(output_path)



# =============================================================================
# 4) M.D AI Business Intelligence — Web Application V1
# =============================================================================

from flask import Flask, request, render_template_string, send_file, jsonify
import csv
import io
import json
import math
import statistics

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50 MB

UPLOAD_DIR = Path(tempfile.gettempdir()) / "md_ai_bi_uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

LAST_RESULTS = []
DATASETS = {}
LAST_DATASET_NAME = None


# -----------------------------------------------------------------------------
# Data Hub
# -----------------------------------------------------------------------------

def _safe_number(value):
    if value is None:
        return None
    s = str(value).strip().replace(",", "").replace("SAR", "").replace("ر.س", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _read_csv(path):
    # utf-8-sig handles Excel-exported CSV files with BOM.
    with open(path, "r", encoding="utf-8-sig", errors="ignore", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        return list(csv.DictReader(f, dialect=dialect))


def _read_excel(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = []
    for i, value in enumerate(rows[0]):
        headers.append(str(value).strip() if value is not None else f"Column_{i+1}")
    data = []
    for row in rows[1:]:
        if not any(v is not None and str(v).strip() for v in row):
            continue
        item = {}
        for i, h in enumerate(headers):
            item[h] = row[i] if i < len(row) else None
        data.append(item)
    return data


def read_dataset(path):
    ext = Path(path).suffix.lower()
    if ext == ".csv":
        return _read_csv(path)
    if ext in (".xlsx", ".xlsm"):
        return _read_excel(path)
    raise ValueError("صيغة البيانات المدعومة حاليًا: CSV / XLSX / XLSM")


def _column(rows, candidates):
    if not rows:
        return None
    keys = list(rows[0].keys())
    normalized = {str(k).strip().lower(): k for k in keys}
    for candidate in candidates:
        if candidate.lower() in normalized:
            return normalized[candidate.lower()]
    for k in keys:
        lk = str(k).lower()
        if any(candidate.lower() in lk for candidate in candidates):
            return k
    return None


def analyze_dataset(rows):
    if not rows:
        return {
            "rows": 0, "columns": 0, "sales": 0, "expenses": 0,
            "profit": 0, "customers": 0, "products": 0,
            "sales_col": None, "expense_col": None,
            "customer_col": None, "product_col": None,
            "date_col": None
        }

    sales_col = _column(rows, [
        "sales", "sale", "revenue", "amount", "total", "net sales",
        "sales amount", "revenue amount", "المبيعات", "الإيرادات", "اجمالي", "المبلغ"
    ])
    expense_col = _column(rows, [
        "expense", "expenses", "cost", "costs", "operating expense",
        "المصروفات", "المصاريف", "التكلفة", "التكاليف"
    ])
    customer_col = _column(rows, [
        "customer", "customer name", "client", "client name", "العميل", "اسم العميل"
    ])
    product_col = _column(rows, [
        "product", "product name", "item", "item name", "sku", "الصنف", "المنتج", "اسم المنتج"
    ])
    date_col = _column(rows, [
        "date", "invoice date", "order date", "transaction date",
        "التاريخ", "تاريخ الفاتورة", "تاريخ الطلب"
    ])

    def total(col):
        if not col:
            return 0.0
        return sum((_safe_number(r.get(col)) or 0) for r in rows)

    sales = total(sales_col)
    expenses = total(expense_col)
    # If an explicit expense column is absent, do not invent expenses.
    profit = sales - expenses if expense_col else 0.0

    customers = len({
        str(r.get(customer_col)).strip()
        for r in rows
        if customer_col and r.get(customer_col) not in (None, "")
    })
    products = len({
        str(r.get(product_col)).strip()
        for r in rows
        if product_col and r.get(product_col) not in (None, "")
    })

    return {
        "rows": len(rows),
        "columns": len(rows[0]),
        "sales": round(sales, 2),
        "expenses": round(expenses, 2),
        "profit": round(profit, 2),
        "customers": customers,
        "products": products,
        "sales_col": sales_col,
        "expense_col": expense_col,
        "customer_col": customer_col,
        "product_col": product_col,
        "date_col": date_col,
    }


def _top_products(rows, limit=8):
    stats = {}
    product_col = _column(rows, [
        "product", "product name", "item", "item name", "sku",
        "الصنف", "المنتج", "اسم المنتج"
    ])
    sales_col = _column(rows, [
        "sales", "sale", "revenue", "amount", "total", "net sales",
        "sales amount", "revenue amount", "المبيعات", "الإيرادات", "اجمالي", "المبلغ"
    ])
    if not product_col or not sales_col:
        return []
    for row in rows:
        name = str(row.get(product_col) or "Unknown").strip()
        amount = _safe_number(row.get(sales_col)) or 0
        stats[name] = stats.get(name, 0) + amount
    return [
        {"name": name, "value": round(value, 2)}
        for name, value in sorted(stats.items(), key=lambda x: x[1], reverse=True)[:limit]
    ]


def _monthly_sales(rows):
    date_col = _column(rows, [
        "date", "invoice date", "order date", "transaction date",
        "التاريخ", "تاريخ الفاتورة", "تاريخ الطلب"
    ])
    sales_col = _column(rows, [
        "sales", "sale", "revenue", "amount", "total", "net sales",
        "sales amount", "revenue amount", "المبيعات", "الإيرادات", "اجمالي", "المبلغ"
    ])
    if not date_col or not sales_col:
        return []

    monthly = {}
    for row in rows:
        raw = str(row.get(date_col) or "").strip()
        if not raw:
            continue
        dt = None
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d-%m-%Y", "%d/%m/%Y",
                    "%m/%d/%Y", "%m-%d-%Y"):
            try:
                dt = datetime.strptime(raw[:10], fmt)
                break
            except ValueError:
                pass
        if not dt:
            continue
        key = dt.strftime("%Y-%m")
        monthly[key] = monthly.get(key, 0) + (_safe_number(row.get(sales_col)) or 0)

    return [{"month": k, "value": round(v, 2)} for k, v in sorted(monthly.items())]


def generate_insights(rows):
    k = analyze_dataset(rows)
    insights = []

    if not rows:
        return ["ارفع ملف CSV أو Excel لبدء التحليل."]

    if k["sales_col"]:
        insights.append(f"تم اكتشاف عمود المبيعات: {k['sales_col']}. إجمالي المبيعات المحسوب = {k['sales']:,.2f}.")
    else:
        insights.append("لم يتم اكتشاف عمود مبيعات واضح. يمكنك إعادة تسمية العمود إلى Sales أو Revenue أو Amount.")

    if k["expense_col"]:
        margin = (k["profit"] / k["sales"] * 100) if k["sales"] else 0
        insights.append(
            f"تم اكتشاف المصروفات. الربح المحسوب = {k['profit']:,.2f}، وهامش الربح التقريبي = {margin:.1f}%."
        )
    else:
        insights.append("لا يوجد عمود مصروفات واضح؛ لذلك لن يخترع النظام رقمًا للمصروفات أو الربح.")

    if k["customer_col"]:
        insights.append(f"عدد العملاء الفريدين المكتشفين: {k['customers']:,}.")
    if k["product_col"]:
        insights.append(f"عدد المنتجات/الأصناف الفريدة: {k['products']:,}.")
    if k["date_col"]:
        monthly = _monthly_sales(rows)
        if len(monthly) >= 2:
            last = monthly[-1]["value"]
            prev = monthly[-2]["value"]
            change = ((last - prev) / prev * 100) if prev else 0
            direction = "ارتفاع" if change >= 0 else "انخفاض"
            insights.append(f"المبيعات الشهرية الأخيرة شهدت {direction} بنسبة {abs(change):.1f}% مقارنة بالشهر السابق.")

    top = _top_products(rows, 1)
    if top:
        insights.append(f"أعلى منتج حسب قيمة المبيعات: {top[0]['name']} بقيمة {top[0]['value']:,.2f}.")

    return insights


# -----------------------------------------------------------------------------
# Document Hub — keep the original engine intact
# -----------------------------------------------------------------------------

def export_to_excel(results, output_path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Document Results"
    headers = [
        "اسم الملف", "رقم المستند", "نوع المستند", "التاريخ",
        "Revision", "Status", "الجهة المرسلة", "دقة الاستخراج"
    ]
    ws.append(headers)
    header_fill = PatternFill(start_color="15243A", end_color="15243A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for r in results:
        ws.append([
            r.source_file, r.document_number, r.document_type, r.date,
            r.revision, r.status, r.sender, r.confidence
        ])
    for col_idx, header in enumerate(headers, start=1):
        max_len = max(
            [len(str(header))]
            + [len(str(ws.cell(row=row, column=col_idx).value or ""))
               for row in range(2, ws.max_row + 1)]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 45)
    ws.freeze_panes = "A2"
    wb.save(output_path)


# -----------------------------------------------------------------------------
# UI
# -----------------------------------------------------------------------------

PAGE_TEMPLATE = r"""
<!doctype html>
<html lang="en" dir="ltr">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>M.D AI Business Intelligence</title>
<script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
<style>
:root{
 --bg:#f5f7fb;--panel:#fff;--navy:#14253d;--navy2:#203b5f;
 --gold:#c8a45d;--text:#182333;--muted:#6b7788;--line:#e7ebf1;
 --green:#21865b;--red:#c54a4a;
}
*{box-sizing:border-box} body{margin:0;background:var(--bg);color:var(--text);
font-family:Inter,"Segoe UI",Tahoma,Arial,sans-serif}
.app{display:flex;min-height:100vh}.side{width:245px;background:var(--navy);color:white;
padding:24px 16px;position:fixed;inset:0 auto 0 0}.brand{padding:4px 12px 24px;border-bottom:1px solid #ffffff18}
.brand b{font-size:21px}.brand span{display:block;color:#b8c5d7;font-size:12px;margin-top:5px}
.nav{margin-top:22px}.nav a{display:block;color:#d7dfeb;text-decoration:none;padding:12px 14px;
border-radius:9px;margin:5px 0;font-size:14px}.nav a:hover,.nav a.active{background:#ffffff12;color:#fff}
.main{margin-left:245px;width:calc(100% - 245px);padding:28px 32px}
.top{display:flex;justify-content:space-between;align-items:center;margin-bottom:24px}.top h1{margin:0;font-size:26px}
.top p{margin:5px 0 0;color:var(--muted);font-size:13px}.pill{background:#fff;border:1px solid var(--line);
padding:9px 13px;border-radius:20px;color:var(--muted);font-size:12px}
.cards{display:grid;grid-template-columns:repeat(4,1fr);gap:16px}.card{background:var(--panel);border:1px solid var(--line);
border-radius:14px;padding:19px;box-shadow:0 3px 14px #14253d0a}.label{font-size:12px;color:var(--muted)}
.value{font-size:25px;font-weight:700;margin-top:8px}.gold{color:#a27c30}.green{color:var(--green)}
.grid{display:grid;grid-template-columns:2fr 1fr;gap:16px;margin-top:16px}.section-title{font-size:15px;font-weight:700;margin-bottom:15px}
.upload{border:1.5px dashed #bdc8d6;text-align:center;padding:30px;border-radius:12px;background:#fafbfd}
input[type=file]{margin:12px 0}.btn{background:var(--navy2);color:#fff;border:0;border-radius:8px;padding:10px 18px;
cursor:pointer;font-weight:600}.btn.goldbtn{background:var(--gold);color:#182333}.btn.secondary{background:#eef2f7;color:var(--navy)}
table{width:100%;border-collapse:collapse;font-size:13px}th,td{text-align:left;padding:10px;border-bottom:1px solid var(--line)}
th{color:var(--muted);font-weight:600}.insight{padding:11px 12px;background:#f7f9fc;border-left:3px solid var(--gold);margin:8px 0;border-radius:6px;font-size:13px;line-height:1.5}
.status{display:inline-block;padding:3px 8px;border-radius:12px;font-size:11px;background:#eef2f7}
.alert{padding:12px 14px;border-radius:9px;margin-bottom:15px;background:#fff4df;color:#79550f;font-size:13px}
.actions{display:flex;gap:9px;flex-wrap:wrap}.small{font-size:12px;color:var(--muted)}
@media(max-width:1000px){.side{width:190px}.main{margin-left:190px;width:calc(100% - 190px)}.cards{grid-template-columns:repeat(2,1fr)}}
@media(max-width:700px){.side{position:static;width:100%;min-height:auto}.app{display:block}.main{margin:0;width:100%;padding:18px}.cards,.grid{grid-template-columns:1fr}}
</style>
</head>
<body>
<div class="app">
<aside class="side">
 <div class="brand"><b>M.D AI</b><span>Business Intelligence Platform</span></div>
 <nav class="nav">
  <a class="active" href="/">▦ Dashboard</a>
  <a href="/data-hub">◈ Data Hub</a>
  <a href="/documents">▤ Documents</a>
  <a href="/analytics">◒ Analytics</a>
  <a href="/ai-analyst">✦ AI Analyst</a>
 </nav>
</aside>
<main class="main">
<div class="top">
 <div><h1>Business Intelligence Dashboard</h1><p>Data-driven insights from your business files</p></div>
 <div class="pill">M.D AI BI • V1</div>
</div>

{% if message %}<div class="alert">{{ message }}</div>{% endif %}

<div class="cards">
 <div class="card"><div class="label">TOTAL SALES</div><div class="value gold">{{ "{:,.2f}".format(k.sales) }}</div></div>
 <div class="card"><div class="label">EXPENSES</div><div class="value">{{ "{:,.2f}".format(k.expenses) }}</div></div>
 <div class="card"><div class="label">CALCULATED PROFIT</div><div class="value green">{{ "{:,.2f}".format(k.profit) }}</div></div>
 <div class="card"><div class="label">DATA ROWS</div><div class="value">{{ "{:,}".format(k.rows) }}</div></div>
</div>

<div class="grid">
 <div class="card">
  <div class="section-title">Sales Trend</div>
  {% if monthly %}
   <canvas id="salesChart" height="110"></canvas>
  {% else %}
   <div class="upload"><b>No time-series data yet</b><div class="small">Upload a dataset containing Date + Sales/Revenue.</div></div>
  {% endif %}
 </div>
 <div class="card">
  <div class="section-title">AI Insights</div>
  {% for i in insights %}<div class="insight">{{ i }}</div>{% endfor %}
 </div>
</div>

<div class="grid">
 <div class="card">
  <div class="section-title">Top Products</div>
  {% if top %}
  <table><tr><th>Product</th><th>Sales</th></tr>
   {% for x in top %}<tr><td>{{ x.name }}</td><td>{{ "{:,.2f}".format(x.value) }}</td></tr>{% endfor %}
  </table>
  {% else %}<div class="small">No Product + Sales columns detected.</div>{% endif %}
 </div>
 <div class="card">
  <div class="section-title">Data Quality</div>
  <div class="insight">Columns detected: {{ k.columns }}</div>
  <div class="insight">Customers: {{ "{:,}".format(k.customers) }}</div>
  <div class="insight">Products: {{ "{:,}".format(k.products) }}</div>
  <div class="small">The engine does not invent missing financial fields.</div>
 </div>
</div>

<div class="card" style="margin-top:16px">
 <div class="section-title">Quick Start</div>
 <div class="actions">
  <a href="/data-hub"><button class="btn goldbtn">Upload Business Data</button></a>
  <a href="/documents"><button class="btn secondary">Analyze Documents</button></a>
  <a href="/ai-analyst"><button class="btn">Ask AI Analyst</button></a>
 </div>
</div>
</main></div>
{% if monthly %}
<script>
new Chart(document.getElementById('salesChart'),{
 type:'line',
 data:{labels:{{ monthly|map(attribute='month')|list|tojson }},
 datasets:[{label:'Sales',data:{{ monthly|map(attribute='value')|list|tojson }},tension:.3}]},
 options:{responsive:true,plugins:{legend:{display:false}},scales:{y:{beginAtZero:true}}}
});
</script>
{% endif %}
</body></html>
"""

DATA_HUB_TEMPLATE = r"""
<!doctype html><html lang="en" dir="ltr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>M.D AI — Data Hub</title><style>
body{font-family:Segoe UI,Arial;background:#f5f7fb;margin:0;color:#182333}.wrap{max-width:950px;margin:45px auto;padding:0 20px}
.card{background:white;border:1px solid #e5e9ef;border-radius:14px;padding:28px;margin-bottom:18px}.btn{background:#203b5f;color:#fff;border:0;padding:11px 20px;border-radius:8px;cursor:pointer}
a{color:#203b5f;text-decoration:none}.hint{color:#697688;font-size:13px;line-height:1.7}.back{display:inline-block;margin-bottom:18px}
</style></head><body><div class="wrap"><a class="back" href="/">← Dashboard</a>
<div class="card"><h1>Data Hub</h1><p class="hint">Upload CSV or Excel business data. The engine detects common Sales, Revenue, Expense, Customer, Product and Date columns.</p>
<form method="POST" enctype="multipart/form-data"><input type="file" name="datafile" accept=".csv,.xlsx,.xlsm" required><button class="btn" type="submit">Load & Analyze</button></form></div>
{% if dataset %}<div class="card"><h2>{{ name }}</h2><p>Rows: {{ k.rows }} • Columns: {{ k.columns }}</p>
<p class="hint">Detected Sales: {{ k.sales_col or "—" }} | Expenses: {{ k.expense_col or "—" }} | Customer: {{ k.customer_col or "—" }} | Product: {{ k.product_col or "—" }} | Date: {{ k.date_col or "—" }}</p>
<a href="/analytics">Open Analytics →</a></div>{% endif %}</div></body></html>
"""

AI_TEMPLATE = r"""
<!doctype html><html lang="en" dir="ltr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>M.D AI — Analyst</title><style>
body{font-family:Segoe UI,Arial;background:#f5f7fb;margin:0;color:#182333}.wrap{max-width:900px;margin:45px auto;padding:0 20px}
.card{background:white;border:1px solid #e5e9ef;border-radius:14px;padding:28px}.btn{background:#203b5f;color:#fff;border:0;padding:11px 20px;border-radius:8px;cursor:pointer}
textarea{width:100%;padding:12px;border:1px solid #d9e0e8;border-radius:8px;min-height:90px;margin:10px 0;box-sizing:border-box}
.answer{background:#f7f9fc;border-left:3px solid #c8a45d;padding:16px;margin-top:18px;line-height:1.7}.muted{color:#697688;font-size:13px}a{color:#203b5f;text-decoration:none}
</style></head><body><div class="wrap"><a href="/">← Dashboard</a><div class="card" style="margin-top:18px">
<h1>✦ AI Analyst</h1><p class="muted">Ask questions about the currently loaded dataset. Calculations come from the Data Engine.</p>
<form method="POST"><textarea name="question" placeholder="Example: What are total sales and the top products?">{{ question or "" }}</textarea><button class="btn">Analyze</button></form>
{% if answer %}<div class="answer">{{ answer|safe }}</div>{% endif %}
</div></div></body></html>
"""

DOCUMENT_TEMPLATE = r"""
<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>M.D AI — Documents</title><style>
body{font-family:Segoe UI,Arial;background:#f5f7fb;margin:0;color:#182333}.wrap{max-width:1100px;margin:35px auto;padding:0 20px}
.card{background:white;border:1px solid #e5e9ef;border-radius:14px;padding:25px;margin-bottom:18px}.btn{background:#203b5f;color:white;border:0;padding:10px 18px;border-radius:8px;cursor:pointer}
table{width:100%;border-collapse:collapse;font-size:13px}th,td{padding:9px;border-bottom:1px solid #eee;text-align:right}th{color:#697688}
.badge{padding:3px 8px;border-radius:10px;background:#eef2f7}.High{background:#dff4e9}.Medium{background:#fff0c9}.Low{background:#fde2e2}a{color:#203b5f;text-decoration:none}
</style></head><body><div class="wrap"><a href="/">← Dashboard</a>
<div class="card"><h1>Document Hub</h1><p>PDF / DOCX / TXT — extraction engine preserved from the original DocuAI application.</p>
<form method="POST" enctype="multipart/form-data"><input type="file" name="files" multiple required><button class="btn">Analyze Documents</button></form></div>
{% if results %}<div class="card"><h2>Results ({{ results|length }})</h2><table>
<tr><th>File</th><th>Document No.</th><th>Type</th><th>Date</th><th>Revision</th><th>Status</th><th>Sender</th><th>Confidence</th></tr>
{% for r in results %}<tr><td>{{r.source_file}}</td><td>{{r.document_number or "—"}}</td><td>{{r.document_type or "—"}}</td><td>{{r.date or "—"}}</td><td>{{r.revision or "—"}}</td><td>{{r.status or "—"}}</td><td>{{r.sender or "—"}}</td><td><span class="badge {{r.confidence}}">{{r.confidence}}</span></td></tr>{% endfor %}
</table><br><a href="/download"><button class="btn">Download Excel</button></a></div>{% endif %}</div></body></html>
"""


def current_rows():
    global LAST_DATASET_NAME
    if LAST_DATASET_NAME and LAST_DATASET_NAME in DATASETS:
        return DATASETS[LAST_DATASET_NAME]
    return []


def analyst_answer(question, rows):
    q = (question or "").lower()
    if not rows:
        return "No business dataset is loaded. Go to Data Hub and upload a CSV or Excel file first."

    k = analyze_dataset(rows)
    top = _top_products(rows)
    monthly = _monthly_sales(rows)
    parts = []

    # Intent-based responses using deterministic calculations.
    if any(x in q for x in ["sales", "revenue", "مبيعات", "إيرادات", "المبيعات"]):
        parts.append(f"<b>Total sales:</b> {k['sales']:,.2f}")
    if any(x in q for x in ["expense", "cost", "مصروف", "تكلفة"]):
        parts.append(
            f"<b>Expenses:</b> {k['expenses']:,.2f}" if k["expense_col"]
            else "<b>Expenses:</b> No expense column was detected."
        )
    if any(x in q for x in ["profit", "ربح", "الربح", "margin", "هامش"]):
        parts.append(
            f"<b>Calculated profit:</b> {k['profit']:,.2f}"
            if k["expense_col"] else
            "<b>Profit:</b> Cannot be calculated reliably because no expense column was detected."
        )
    if any(x in q for x in ["product", "item", "منتج", "صنف"]):
        if top:
            parts.append("<b>Top products:</b><br>" + "<br>".join(
                f"{i+1}. {x['name']} — {x['value']:,.2f}" for i, x in enumerate(top[:5])
            ))
        else:
            parts.append("No Product/Item + Sales columns were detected.")

    if any(x in q for x in ["customer", "client", "عميل", "عملاء"]):
        parts.append(
            f"<b>Unique customers:</b> {k['customers']:,}"
            if k["customer_col"] else
            "No Customer/Client column was detected."
        )

    if any(x in q for x in ["trend", "monthly", "month", "اتجاه", "شهري", "شهر"]):
        if monthly:
            parts.append("<b>Monthly sales:</b><br>" + "<br>".join(
                f"{x['month']}: {x['value']:,.2f}" for x in monthly[-12:]
            ))
        else:
            parts.append("No usable Date + Sales columns were detected.")

    if not parts:
        parts = generate_insights(rows)

    return "<br><br>".join(parts)


@app.route("/", methods=["GET"])
def dashboard():
    rows = current_rows()
    k = analyze_dataset(rows)
    return render_template_string(
        PAGE_TEMPLATE,
        k=k, monthly=_monthly_sales(rows), top=_top_products(rows),
        insights=generate_insights(rows), message=None
    )


@app.route("/data-hub", methods=["GET", "POST"])
def data_hub():
    global LAST_DATASET_NAME
    message = None
    if request.method == "POST":
        f = request.files.get("datafile")
        if not f or not f.filename:
            message = "اختر ملف CSV أو Excel."
        else:
            ext = Path(f.filename).suffix.lower()
            if ext not in (".csv", ".xlsx", ".xlsm"):
                message = "صيغة غير مدعومة. استخدم CSV أو XLSX أو XLSM."
            else:
                save_path = UPLOAD_DIR / Path(f.filename).name
                f.save(save_path)
                try:
                    rows = read_dataset(str(save_path))
                    DATASETS[f.filename] = rows
                    LAST_DATASET_NAME = f.filename
                except Exception as e:
                    message = f"تعذر قراءة الملف: {e}"

    rows = current_rows()
    return render_template_string(
        DATA_HUB_TEMPLATE, dataset=bool(rows), name=LAST_DATASET_NAME,
        k=analyze_dataset(rows), message=message
    )


@app.route("/analytics")
def analytics():
    return dashboard()


@app.route("/ai-analyst", methods=["GET", "POST"])
def ai_analyst():
    question = ""
    answer = None
    if request.method == "POST":
        question = request.form.get("question", "")
        answer = analyst_answer(question, current_rows())
    return render_template_string(AI_TEMPLATE, question=question, answer=answer)


@app.route("/documents", methods=["GET", "POST"])
def documents():
    global LAST_RESULTS
    results = []
    if request.method == "POST":
        files = request.files.getlist("files")
        for f in files:
            if not f.filename:
                continue
            filename = Path(f.filename).name
            save_path = UPLOAD_DIR / filename
            f.save(save_path)
            try:
                text = read_any(str(save_path))
                extracted = extract_fields(text, source_file=filename)
                results.append(extracted)
            except Exception as e:
                print(f"خطأ في معالجة {filename}: {e}")
        LAST_RESULTS = results
    else:
        results = LAST_RESULTS
    return render_template_string(DOCUMENT_TEMPLATE, results=results)


@app.route("/download")
def download():
    if not LAST_RESULTS:
        return "لا توجد نتائج بعد. الرجاء تحليل مستندات أولاً.", 400
    output_path = UPLOAD_DIR / "md_ai_document_results.xlsx"
    export_to_excel(LAST_RESULTS, str(output_path))
    return send_file(
        output_path, as_attachment=True,
        download_name="md_ai_document_results.xlsx"
    )


@app.route("/api/dashboard")
def api_dashboard():
    rows = current_rows()
    return jsonify({
        "kpis": analyze_dataset(rows),
        "monthly_sales": _monthly_sales(rows),
        "top_products": _top_products(rows),
        "insights": generate_insights(rows),
        "dataset": LAST_DATASET_NAME
    })


# -----------------------------------------------------------------------------
# Local / Render entry point
# -----------------------------------------------------------------------------

def find_free_port(preferred=5000):
    for port in [preferred] + list(range(5001, 5011)):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            if s.connect_ex(("127.0.0.1", port)) != 0:
                return port
    return preferred


def open_browser(url):
    threading.Timer(1.2, lambda: webbrowser.open(url)).start()


if __name__ == "__main__":
    env_port = os.environ.get("PORT")
    is_cloud = env_port is not None
    port = int(env_port) if is_cloud else find_free_port(5000)
    host = "0.0.0.0" if is_cloud else "127.0.0.1"

    print("=" * 60)
    print("  M.D AI Business Intelligence V1 جاهز للعمل")
    print("  Data Hub + Analytics Engine + AI Analyst + Document Hub")
    print("=" * 60)

    if not is_cloud:
        open_browser(f"http://{host}:{port}")

    app.run(debug=False, host=host, port=port)
