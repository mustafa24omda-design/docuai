"""
DocuAI - Intelligent Project Document Management System
=========================================================
نسخة مدمجة بالكامل في ملف واحد (app.py) — بدون الحاجة لمجلدات فرعية،
لتسهيل الرفع عبر واجهة GitHub الويب والنشر المباشر على Render.

يحتوي هذا الملف على:
1. محرك الاستخراج (Extraction Engine) بالـ Regex.
2. قارئات الملفات (PDF / DOCX / TXT).
3. واجهة ويب (Flask) لرفع المستندات ومعاينة/تنزيل النتائج.
"""

import os
import re
import socket
import tempfile
import threading
import webbrowser
from pathlib import Path
from datetime import datetime
from dataclasses import dataclass, asdict
from typing import Optional

from flask import Flask, request, render_template_string, send_file

# OCR is loaded lazily so the app can still start when OCR dependencies are missing.


# =============================================================================
# 1) محرك الاستخراج (Extraction Engine)
# =============================================================================

DOC_NUMBER_PATTERNS = [
    r"(?:Document\s*(?:No\.?|Number|#)\s*[:\-]?\s*)([A-Z0-9][A-Z0-9\-\/\.]{3,30})",
    r"(?:Doc\.?\s*(?:No\.?|#)\s*[:\-]?\s*)([A-Z0-9][A-Z0-9\-\/\.]{3,30})",
    r"(?:Ref(?:erence)?\.?\s*(?:No\.?)?\s*[:\-]?\s*)([A-Z0-9][A-Z0-9\-\/\.]{3,30})",
    r"\b([A-Z]{2,6}-[A-Z0-9]{2,10}-\d{2,6}(?:-[A-Z0-9]{1,4})?)\b",
]

DOC_TYPE_KEYWORDS = {
    "Transmittal": ["transmittal", "خطاب إحالة", "إحالة مستند"],
    "RFI": ["request for information", "rfi", "طلب معلومات"],
    "Submittal": ["submittal", "material submittal", "shop drawing submittal"],
    "Shop Drawing": ["shop drawing", "مخطط تنفيذي"],
    "Method Statement": ["method statement", "منهجية تنفيذ"],
    "Inspection Request": ["inspection request", "ir ", "طلب معاينة"],
    "Non-Conformance Report": ["non-conformance", "ncr", "تقرير عدم مطابقة"],
    "Minutes of Meeting": ["minutes of meeting", "mom", "محضر اجتماع"],
    "Design Drawing": ["design drawing", "مخطط تصميمي"],
    "Correspondence": ["letter", "correspondence", "خطاب", "مراسلة"],
    "Progress Report": ["progress report", "تقرير سير العمل"],
    "Invoice": ["invoice", "فاتورة"],
}

DATE_PATTERNS = [
    r"\b(\d{4}[-/]\d{1,2}[-/]\d{1,2})\b",
    r"\b(\d{1,2}[-/]\d{1,2}[-/]\d{4})\b",
    r"\b(\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4})\b",
    r"\b([A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4})\b",
]

DATE_FORMATS = [
    "%Y-%m-%d", "%Y/%m/%d",
    "%d-%m-%Y", "%d/%m/%Y",
    "%d %B %Y", "%d %b %Y",
    "%B %d, %Y", "%b %d, %Y", "%B %d %Y", "%b %d %Y",
]

REVISION_PATTERNS = [
    r"(?:Revision|Rev\.?)\s*[:\-]?\s*([A-Z]?\d{1,3}|[A-Z])\b",
    r"(?:مراجعة|مراجعه|النسخة)\s*[:\-]?\s*(\d{1,3})",
]

STATUS_KEYWORDS = {
    "Approved": ["approved", "معتمد"],
    "Approved as Noted": ["approved as noted", "معتمد مع ملاحظات"],
    "Rejected": ["rejected", "not approved", "مرفوض"],
    "Revise and Resubmit": ["revise and resubmit", "أعد التقديم"],
    "For Review": ["for review", "under review", "قيد المراجعة"],
    "For Approval": ["for approval", "للاعتماد"],
    "For Information": ["for information", "for your information", "للعلم"],
    "Closed": ["closed", "مغلق"],
    "Open": ["open", "مفتوح"],
    "Pending": ["pending", "معلق"],
}

SENDER_PATTERNS = [
    r"(?:From|Sender|Originator|Issued\s*By)\s*[:\-]\s*([A-Za-z0-9&.,\-\s]{2,60})(?:\n|$)",
    r"(?:من|المرسل|جهة الإصدار)\s*[:\-]\s*([\u0600-\u06FF0-9A-Za-z&.,\-\s]{2,60})(?:\n|$)",
]


@dataclass
class ExtractedDocument:
    source_file: str
    document_number: Optional[str] = None
    document_type: Optional[str] = None
    date: Optional[str] = None
    revision: Optional[str] = None
    status: Optional[str] = None
    sender: Optional[str] = None
    confidence: Optional[str] = None

    def to_dict(self):
        return asdict(self)


def _first_match(patterns, text, flags=re.IGNORECASE):
    for pat in patterns:
        m = re.search(pat, text, flags)
        if m:
            return m.group(1).strip().rstrip(".,;")
    return None


def _extract_document_type(text: str) -> Optional[str]:
    lowered = text.lower()
    for doc_type, keywords in DOC_TYPE_KEYWORDS.items():
        for kw in keywords:
            if kw.lower() in lowered:
                return doc_type
    return None


def _normalize_date(raw_date: str) -> Optional[str]:
    if not raw_date:
        return None
    cleaned = raw_date.strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(cleaned, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return cleaned


def _match_status_in(snippet: str) -> Optional[str]:
    lowered = snippet.lower()
    all_kw = [(status, kw) for status, kws in STATUS_KEYWORDS.items() for kw in kws]
    all_kw.sort(key=lambda x: len(x[1]), reverse=True)
    for status, kw in all_kw:
        if kw.lower() in lowered:
            return status
    return None


def _extract_status(text: str) -> Optional[str]:
    label_match = re.search(r"(?:Status|الحالة)\s*[:\-]\s*(.+)", text, re.IGNORECASE)
    if label_match:
        found = _match_status_in(label_match.group(1))
        if found:
            return found
    return _match_status_in(text)


def extract_fields(text: str, source_file: str = "") -> ExtractedDocument:
    doc_number = _first_match(DOC_NUMBER_PATTERNS, text)
    doc_type = _extract_document_type(text)
    raw_date = _first_match(DATE_PATTERNS, text)
    date = _normalize_date(raw_date) if raw_date else None
    revision = _first_match(REVISION_PATTERNS, text)
    status = _extract_status(text)
    sender = _first_match(SENDER_PATTERNS, text)

    found_count = sum(1 for v in [doc_number, doc_type, date, revision, status, sender] if v)
    if found_count >= 5:
        confidence = "High"
    elif found_count >= 3:
        confidence = "Medium"
    else:
        confidence = "Low"

    return ExtractedDocument(
        source_file=source_file,
        document_number=doc_number,
        document_type=doc_type,
        date=date,
        revision=revision,
        status=status,
        sender=sender,
        confidence=confidence,
    )


# =============================================================================
# 2) قراءة الملفات (PDF / DOCX / TXT)
# =============================================================================

def read_pdf(path: str) -> str:
    """
    Read a PDF. First try normal text extraction. If a page has little/no
    selectable text, OCR the page using Tesseract.

    OCR requirements:
      pip install pytesseract pdf2image pillow
    Also install the Tesseract OCR engine on the operating system.
    """
    import pdfplumber

    text_parts = []
    pages_needing_ocr = []

    # 1) Fast path: extract selectable PDF text.
    with pdfplumber.open(path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            extracted = (page.extract_text() or "").strip()
            text_parts.append(extracted)
            if len(extracted) < 20:
                pages_needing_ocr.append(page_no)

    # 2) OCR only pages that appear scanned/image-only.
    if pages_needing_ocr:
        try:
            import pytesseract
            from pdf2image import convert_from_path
        except ImportError as exc:
            raise RuntimeError(
                "هذا الـPDF يحتوي صفحات ممسوحة ضوئياً وتحتاج OCR. "
                "ثبّت: pip install pytesseract pdf2image pillow"
            ) from exc

        try:
            images = convert_from_path(
                path,
                dpi=250,
                first_page=min(pages_needing_ocr),
                last_page=max(pages_needing_ocr)
            )
        except Exception as exc:
            raise RuntimeError(
                "تعذر تحويل صفحات PDF إلى صور. تأكد من تثبيت Poppler "
                "وإضافته إلى PATH في Windows."
            ) from exc

        for index, page_no in enumerate(range(min(pages_needing_ocr), max(pages_needing_ocr) + 1)):
            if page_no not in pages_needing_ocr:
                continue

            image = images[index]
            try:
                # Arabic + English OCR. If ara is unavailable, fall back to English.
                langs = "eng+ara"
                try:
                    ocr_text = pytesseract.image_to_string(image, lang=langs)
                except Exception:
                    ocr_text = pytesseract.image_to_string(image, lang="eng")

                text_parts[page_no - 1] = ocr_text.strip()
            except Exception as exc:
                raise RuntimeError(
                    f"فشل OCR في الصفحة {page_no}: {exc}"
                ) from exc

    return "\n".join(text_parts)


def read_docx(path: str) -> str:
    import docx
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def read_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


READERS = {".pdf": read_pdf, ".docx": read_docx, ".txt": read_txt}


def read_any(path: str) -> str:
    ext = Path(path).suffix.lower()
    reader = READERS.get(ext)
    if reader is None:
        raise ValueError(f"صيغة الملف غير مدعومة: {ext}")
    return reader(path)


# =============================================================================
# 3) تصدير النتائج إلى Excel
# =============================================================================

def export_to_excel(results, output_path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "DocuAI Results"

    headers = [
        "اسم الملف", "رقم المستند", "نوع المستند",
        "التاريخ", "Revision", "Status", "الجهة المرسلة", "دقة الاستخراج",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for r in results:
        ws.append([
            r.source_file, r.document_number, r.document_type,
            r.date, r.revision, r.status, r.sender, r.confidence,
        ])

    for col_idx, header in enumerate(headers, start=1):
        max_len = max(
            [len(str(header))]
            + [len(str(ws.cell(row=row, column=col_idx).value or "")) for row in range(2, ws.max_row + 1)]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 45)

    ws.freeze_panes = "A2"
    wb.save(output_path)


# =============================================================================
# 4) واجهة الويب (Flask)
# =============================================================================

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # 25 MB

UPLOAD_DIR = Path(tempfile.gettempdir()) / "docuai_uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

LAST_RESULTS = []

PAGE_TEMPLATE = """
<!DOCTYPE html>
<html lang="ar" dir="rtl">
<head>
<meta charset="UTF-8">
<title>DocuAI - نظام إدارة مستندات المشاريع الذكي</title>
<style>
  body { font-family: 'Segoe UI', Tahoma, sans-serif; background:#f4f6f9; margin:0; padding:0; }
  header { background:#1f4e78; color:#fff; padding:24px 40px; }
  header h1 { margin:0; font-size:24px; }
  header p { margin:4px 0 0; opacity:.85; font-size:14px; }
  .container { max-width:1000px; margin:30px auto; padding:0 20px; }
  .card { background:#fff; border-radius:10px; padding:24px; box-shadow:0 2px 8px rgba(0,0,0,.06); margin-bottom:24px; }
  input[type=file] { display:block; margin:12px 0; }
  button { background:#1f4e78; color:#fff; border:none; padding:10px 22px; border-radius:6px; cursor:pointer; font-size:15px; }
  button:hover { background:#163a5c; }
  table { width:100%; border-collapse:collapse; margin-top:16px; }
  th, td { padding:10px 12px; border-bottom:1px solid #eee; text-align:right; font-size:14px; }
  th { background:#f0f3f7; color:#1f4e78; }
  .badge { padding:3px 10px; border-radius:12px; font-size:12px; color:#fff; }
  .High { background:#2e7d32; } .Medium { background:#f9a825; } .Low { background:#c62828; }
  footer { text-align:center; color:#888; font-size:12px; margin:30px 0; }
</style>
</head>
<body>
<header>
  <h1>📄 DocuAI — Intelligent Project Document Management System</h1>
  <p>استخراج تلقائي لرقم المستند، النوع، التاريخ، Revision، Status، والجهة المرسلة — مع OCR للملفات الممسوحة</p>
</header>
<div class="container">
  <div class="card">
    <form method="POST" enctype="multipart/form-data">
      <label>ارفع ملف أو أكثر (PDF / DOCX / TXT) — يدعم OCR للمستندات المصورة:</label>
      <input type="file" name="files" multiple required>
      <button type="submit">🔍 تحليل المستندات + OCR</button>
    </form>
  </div>

  {% if results %}
  <div class="card">
    <h3>نتائج التحليل ({{ results|length }} مستند)</h3>
    <table>
      <tr>
        <th>اسم الملف</th><th>رقم المستند</th><th>النوع</th><th>التاريخ</th>
        <th>Revision</th><th>Status</th><th>الجهة المرسلة</th><th>الثقة</th>
      </tr>
      {% for r in results %}
      <tr>
        <td>{{ r.source_file }}</td>
        <td>{{ r.document_number or '—' }}</td>
        <td>{{ r.document_type or '—' }}</td>
        <td>{{ r.date or '—' }}</td>
        <td>{{ r.revision or '—' }}</td>
        <td>{{ r.status or '—' }}</td>
        <td>{{ r.sender or '—' }}</td>
        <td><span class="badge {{ r.confidence }}">{{ r.confidence }}</span></td>
      </tr>
      {% endfor %}
    </table>
    <form method="GET" action="/download">
      <button type="submit">⬇️ تنزيل النتائج (Excel)</button>
    </form>
  </div>
  {% endif %}
</div>
<footer>DocuAI © 2026</footer>
</body>
</html>
"""


@app.route("/", methods=["GET", "POST"])
def index():
    global LAST_RESULTS
    results = []
    if request.method == "POST":
        files = request.files.getlist("files")
        for f in files:
            if not f.filename:
                continue
            save_path = UPLOAD_DIR / f.filename
            f.save(save_path)
            try:
                text = read_any(str(save_path))
                extracted = extract_fields(text, source_file=f.filename)
                results.append(extracted)
            except Exception as e:
                print(f"خطأ في معالجة {f.filename}: {e}")
        LAST_RESULTS = results

    return render_template_string(PAGE_TEMPLATE, results=results)


@app.route("/download")
def download():
    if not LAST_RESULTS:
        return "لا توجد نتائج بعد. الرجاء رفع مستندات أولاً.", 400
    output_path = UPLOAD_DIR / "docuai_results.xlsx"
    export_to_excel(LAST_RESULTS, str(output_path))
    return send_file(output_path, as_attachment=True, download_name="docuai_results.xlsx")


def find_free_port(preferred=5000):
    for port in [preferred] + list(range(5001, 5011)):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            if s.connect_ex(("127.0.0.1", port)) != 0:
                return port
    return preferred


def open_browser(url):
    threading.Timer(1.2, lambda: webbrowser.open(url)).start()


if __name__ == "__main__":
    env_port = os.environ.get("PORT")
    is_cloud = env_port is not None
    port = int(env_port) if is_cloud else find_free_port(5000)
    host = "0.0.0.0" if is_cloud else "127.0.0.1"
    url = f"http://{host}:{port}"

    print("=" * 60)
    print("  DocuAI جاهز للعمل ✅")
    if not is_cloud:
        print(f"  افتح المتصفح على: {url}")
        print("  لإيقاف البرنامج اضغط CTRL+C")
    print("=" * 60)

    if not is_cloud:
        open_browser(url)

    app.run(debug=False, host=host, port=port)
